"""Render everything PAIR knows about a candidate as a résumé JobDiva can store.

Why this exists (2026-09-25). Every JobDiva profile Launch PAIR created for a
person JobDiva did not already have (LinkedIn-Unipile / LinkedIn-Exa /
LinkedIn-DeepSearch ...) came out blank. A read-back of PAIR-created profiles
(``bi/CandidatesProfileDetail`` + ``bi/ResumeDetail``) showed a 0-byte ``.txt``
résumé whose PLAINTEXT was 3 characters, JobDiva's ``Auto_…`` placeholder
email, and empty phone / city / state / zip / experience / education. Only the
name, patched on afterwards, survived. Two causes:

1. ``CreateJobApplicationWithResume`` reads the résumé from ``filecontent``
   (the base64 file; ``.pdf`` / ``.docx`` / ``.rtf`` / ``.txt``). ``textfile``
   is only the "Alternate text resume". PAIR always sent ``filecontent: ""``,
   so JobDiva had no résumé to store or parse.
2. For a LinkedIn row the text PAIR did send was ``NAME / Email | Phone /
   (Profile sourced via PAIR)``: LinkedIn rows carry their profile as
   structured data (experience, education, skills, summary), never as
   ``resume_text``, and nothing rendered it.

``build_profile_resume`` renders the résumé (the row's own résumé when it has a
real one, else a LinkedIn-style résumé from the structured profile),
``resume_to_docx`` turns it into the Word file services/jobdiva.py uploads, and
``normalize_linkedin_profile`` keeps every section of a Unipile profile payload
so sourcing can carry the full profile through to Launch PAIR. Everything here
is pure (no I/O) apart from building the document in memory.
"""
from __future__ import annotations

import html
import io
import json
import logging
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Any, Dict, Iterable, List, Optional, Tuple
from urllib.parse import unquote, urlparse

logger = logging.getLogger(__name__)

# Bounds keep one candidate small on the Step-5 stream and in
# sourced_candidates.data; a real LinkedIn profile rarely reaches them.
MAX_EXPERIENCE = 20
MAX_EDUCATION = 10
MAX_SKILLS = 80
MAX_LIST = 15  # certifications / languages / projects / volunteering / links
MAX_FIELD_CHARS = 3000  # summary, one role's description
MAX_PROFILE_TEXT_CHARS = 8000

_MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")

# Section headings the renderer emits. resume_to_docx styles exactly these.
SECTION_SUMMARY = "PROFESSIONAL SUMMARY"
SECTION_EXPERIENCE = "PROFESSIONAL EXPERIENCE"
SECTION_EDUCATION = "EDUCATION"
SECTION_SKILLS = "SKILLS"
SECTION_CERTIFICATIONS = "CERTIFICATIONS"
SECTION_LANGUAGES = "LANGUAGES"
SECTION_PROJECTS = "PROJECTS"
SECTION_VOLUNTEERING = "VOLUNTEER EXPERIENCE"
SECTION_PROFILE_TEXT = "LINKEDIN PROFILE"
SECTION_TITLES = frozenset({
    SECTION_SUMMARY, SECTION_EXPERIENCE, SECTION_EDUCATION, SECTION_SKILLS,
    SECTION_CERTIFICATIONS, SECTION_LANGUAGES, SECTION_PROJECTS,
    SECTION_VOLUNTEERING, SECTION_PROFILE_TEXT,
})

# Body sections that describe the person's career. A résumé with none of them
# is "thin": worth re-reading the LinkedIn profile before creating anything.
_CAREER_SECTIONS = frozenset({"experience", "summary", "resume", "profile_text"})

# Names the sourcing pipeline falls back to when it has none
# (services/unipile.py `_resolve_candidate_name`, the Step-5 save payload).
_PLACEHOLDER_NAMES = frozenset({
    "", "unknown", "unknown unknown", "unknown candidate", "unnamed candidate",
    "linkedin candidate", "professional candidate", "candidate", "n/a", "na",
})
_HONORIFICS = frozenset({"dr", "dr.", "mr", "mr.", "mrs", "mrs.", "ms", "ms.", "miss", "prof", "prof."})

# Lines Exa's LinkedIn crawl carries that say nothing about the person.
_PROFILE_TEXT_NOISE_RE = re.compile(
    r"^\s*(?:\d[\d,.]*\s*\+?\s*(?:connections?|followers?)\b.*|total experience:.*)$",
    re.IGNORECASE | re.MULTILINE,
)
# Characters XML 1.0 (and therefore .docx) cannot carry.
_XML_INVALID_RE = re.compile("[^\u0009\u000A\u000D -퟿-�\U00010000-\U0010FFFF]")


# ---------------------------------------------------------------------------
# Small coercion helpers
# ---------------------------------------------------------------------------

def _as_dict(value: Any) -> Dict[str, Any]:
    if isinstance(value, dict):
        return value
    if isinstance(value, str) and value.strip().startswith("{"):
        try:
            parsed = json.loads(value)
        except (TypeError, ValueError):
            return {}
        return parsed if isinstance(parsed, dict) else {}
    return {}


def _as_list(value: Any) -> List[Any]:
    if isinstance(value, list):
        return value
    if isinstance(value, str) and value.strip().startswith("["):
        try:
            parsed = json.loads(value)
        except (TypeError, ValueError):
            return []
        return parsed if isinstance(parsed, list) else []
    return []


def _text(value: Any, limit: int = 300) -> str:
    """Whitespace-normalised text, newlines kept (descriptions have paragraphs)."""
    if value is None or isinstance(value, (dict, list)):
        return ""
    # Exa's crawl keeps HTML entities ("Planning &amp; Construction").
    text = html.unescape(str(value)).replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) > limit:
        text = text[:limit].rstrip()
    return text


def _line(value: Any, limit: int = 300) -> str:
    """Single-line text."""
    return re.sub(r"\s+", " ", _text(value, limit)).strip()


def _date_text(value: Any) -> str:
    """Display form ("Jan 2020" / "2020") of the date shapes LinkedIn payloads use:
    {"year", "month"} dicts, "1/2020", ISO strings, bare years, epoch millis."""
    if value is None or value == "" or value is False:
        return ""
    if isinstance(value, dict):
        try:
            year = int(value.get("year"))
        except (TypeError, ValueError):
            return ""
        try:
            month = int(value.get("month") or 0)
        except (TypeError, ValueError):
            month = 0
        return f"{_MONTHS[month - 1]} {year}" if 1 <= month <= 12 else str(year)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = int(value)
        if 1900 <= number <= 2100:
            return str(number)
        if number > 10 ** 11:  # epoch milliseconds
            try:
                stamp = datetime.fromtimestamp(number / 1000, tz=timezone.utc)
            except (OverflowError, OSError, ValueError):
                return ""
            return f"{_MONTHS[stamp.month - 1]} {stamp.year}"
        return ""
    text = _line(value, 40)
    match = re.match(r"^(\d{4})-(\d{1,2})(?:-\d{1,2})?(?:[T ].*)?$", text)
    if match and 1 <= int(match.group(2)) <= 12:
        return f"{_MONTHS[int(match.group(2)) - 1]} {match.group(1)}"
    match = re.match(r"^(\d{1,2})/(?:\d{1,2}/)?(\d{4})$", text)
    if match and 1 <= int(match.group(1)) <= 12:
        return f"{_MONTHS[int(match.group(1)) - 1]} {match.group(2)}"
    return text


def _date_range(start: str, end: str, current: bool = False) -> str:
    if start and (end or current):
        return f"{start} - {end or 'Present'}"
    if start:
        # LinkedIn leaves the end off a role that is still going.
        return f"{start} - Present"
    return end


def _names(value: Any, limit: int) -> List[str]:
    """Distinct names from a list of strings / {"name"} / {"skill"} dicts."""
    out: List[str] = []
    seen = set()
    for item in _as_list(value):
        if isinstance(item, dict):
            item = item.get("name") or item.get("skill") or item.get("skill_name") or item.get("title")
        name = _line(item, 120)
        key = name.casefold()
        if name and key not in seen:
            seen.add(key)
            out.append(name)
        if len(out) >= limit:
            break
    return out


def _compact(item: Dict[str, Any]) -> Dict[str, Any]:
    return {k: v for k, v in item.items() if v not in ("", None, [], {}, False)}


def _org_name(value: Any) -> str:
    if isinstance(value, dict):
        value = value.get("name") or value.get("title")
    return _line(value, 200)


def public_linkedin_url(*candidates: Any) -> str:
    """First public ``linkedin.com/in/<slug>`` URL among ``candidates`` ("" if none).

    Recruiter deep links (``/talent/...``) need an RPS seat and are never a
    candidate's public identity, so they are skipped.
    """
    for value in candidates:
        url = _line(value, 400)
        if "linkedin.com/in/" not in url.lower():
            continue
        if not url.lower().startswith(("http://", "https://")):
            url = "https://" + url.lstrip("/")
        return url
    return ""


# JobDiva's social-network types (GET /apiv2/jobdiva/getSocialNetworkTypes,
# 2026-09-25): "Professional Website", "MySpace", "LinkedIn", "X", "Facebook",
# "YouTube", "StackOverflow", "Instagram", "GitHub". Matched on the URL's host.
_SOCIAL_NETWORK_HOSTS = (
    ("linkedin.com", "LinkedIn"),
    ("github.com", "GitHub"),
    ("stackoverflow.com", "StackOverflow"),
    ("twitter.com", "X"),
    ("x.com", "X"),
    ("facebook.com", "Facebook"),
    ("youtube.com", "YouTube"),
    ("instagram.com", "Instagram"),
)
PROFESSIONAL_WEBSITE = "Professional Website"


def _as_url(value: Any) -> str:
    url = _line(value, 400)
    if not url or " " in url or "." not in url or "@" in url.split("/")[0]:
        return ""
    if not url.lower().startswith(("http://", "https://")):
        url = "https://" + url.lstrip("/")
    return url


def jobdiva_social_links(row: Dict[str, Any], data: Optional[Dict[str, Any]] = None) -> Dict[str, str]:
    """JobDiva social-network name -> URL for every link the row carries: the
    public LinkedIn profile, plus GitHub / portfolio / websites from the LinkedIn
    profile and the résumé parse. Recruiter-only LinkedIn links are skipped."""
    data = data if isinstance(data, dict) else _as_dict(row.get("data"))
    profile = _as_dict(data.get("linkedin_profile"))
    enhanced = _as_dict(data.get("enhanced_info")) or _as_dict(row.get("enhanced_info"))
    urls = _as_dict(data.get("urls")) or _as_dict(enhanced.get("urls"))
    out: Dict[str, str] = {}
    linkedin = public_linkedin_url(
        row.get("profile_url"), profile.get("public_profile_url"), urls.get("linkedin"), urls.get("linkedin_url"),
    )
    if linkedin:
        out["LinkedIn"] = linkedin
    for value in [*_as_list(profile.get("websites")), urls.get("github"), urls.get("portfolio"), urls.get("website")]:
        url = _as_url(value)
        if not url:
            continue
        host = urlparse(url).netloc.lower().split(":")[0]
        name = next(
            (label for domain, label in _SOCIAL_NETWORK_HOSTS if host == domain or host.endswith("." + domain)),
            PROFESSIONAL_WEBSITE,
        )
        if name != "LinkedIn":  # only the public /in/ profile, taken above
            out.setdefault(name, url)
    return out


def alternate_email_of(row: Dict[str, Any], data: Optional[Dict[str, Any]], primary: str) -> str:
    """A second real email for the person (enrichment providers, LinkedIn, the
    résumé parse) that differs from ``primary``; "" when there is none."""
    from utils.email_utils import is_placeholder_email

    data = data if isinstance(data, dict) else _as_dict(row.get("data"))
    enrichment = _as_dict(data.get("zoominfo_contact_enrichment"))
    profile = _as_dict(data.get("linkedin_profile"))
    enhanced = _as_dict(data.get("enhanced_info")) or _as_dict(row.get("enhanced_info"))
    primary_key = (primary or "").strip().lower()
    for value in (
        enrichment.get("workEmail"), enrichment.get("personalEmail"),
        *_as_list(profile.get("emails")), enhanced.get("email"), row.get("email"),
    ):
        email = _line(value, 200)
        if email and "@" in email and not is_placeholder_email(email) and email.lower() != primary_key:
            return email
    return ""


def linkedin_public_identifier(url: Any) -> str:
    """``https://www.linkedin.com/in/jane-doe-123/`` -> ``jane-doe-123``."""
    public = public_linkedin_url(url)
    if not public:
        return ""
    parts = [p for p in urlparse(public).path.split("/") if p]
    if "in" in parts and parts.index("in") + 1 < len(parts):
        slug = unquote(parts[parts.index("in") + 1]).strip()
        if slug and " " not in slug:
            return slug
    return ""


# ---------------------------------------------------------------------------
# LinkedIn profile normalisation (Unipile /users/{id}?linkedin_sections=*)
# ---------------------------------------------------------------------------

def _experience_item(raw: Any) -> Optional[Dict[str, Any]]:
    if not isinstance(raw, dict):
        return None
    title = _line(raw.get("position") or raw.get("title") or raw.get("job_title") or raw.get("role"), 200)
    company = _org_name(raw.get("company") or raw.get("company_name") or raw.get("organization"))
    if not (title or company):
        return None
    return _compact({
        "title": title,
        "company": company,
        "location": _line(raw.get("location"), 200),
        "start": _date_text(raw.get("start") or raw.get("start_date") or raw.get("startDate")),
        "end": _date_text(raw.get("end") or raw.get("end_date") or raw.get("endDate")),
        "current": bool(raw.get("current") or raw.get("is_current")),
        "description": _text(raw.get("description") or raw.get("summary"), MAX_FIELD_CHARS),
        "skills": _names(raw.get("skills"), 20),
    })


def _education_item(raw: Any) -> Optional[Dict[str, Any]]:
    if not isinstance(raw, dict):
        return None
    school = _org_name(raw.get("school") or raw.get("institution") or raw.get("school_name") or raw.get("schoolName"))
    degree = _line(raw.get("degree") or raw.get("degree_name") or raw.get("degreeName"), 200)
    field = _line(raw.get("field_of_study") or raw.get("fieldOfStudy") or raw.get("field") or raw.get("major"), 200)
    if not (school or degree or field):
        return None
    return _compact({
        "school": school,
        "degree": degree,
        "field_of_study": field,
        "start": _date_text(raw.get("start") or raw.get("start_date")),
        "end": _date_text(raw.get("end") or raw.get("end_date") or raw.get("year")),
        "description": _text(raw.get("description") or raw.get("activities"), 600),
    })


def _certification_item(raw: Any) -> Optional[Dict[str, Any]]:
    if isinstance(raw, str):
        raw = {"name": raw}
    if not isinstance(raw, dict):
        return None
    name = _line(raw.get("name") or raw.get("certification_name") or raw.get("title"), 200)
    if not name:
        return None
    return _compact({
        "name": name,
        "issuer": _org_name(raw.get("organization") or raw.get("authority") or raw.get("issuer")),
        "date": _date_text(raw.get("issue_date") or raw.get("start") or raw.get("date") or raw.get("year")),
    })


def _language_item(raw: Any) -> Optional[Dict[str, Any]]:
    if isinstance(raw, str):
        raw = {"name": raw}
    if not isinstance(raw, dict):
        return None
    name = _line(raw.get("name") or raw.get("language"), 80)
    if not name:
        return None
    return _compact({"name": name, "proficiency": _line(raw.get("proficiency") or raw.get("level"), 80)})


def _project_item(raw: Any) -> Optional[Dict[str, Any]]:
    if not isinstance(raw, dict):
        return None
    name = _line(raw.get("name") or raw.get("title"), 200)
    if not name:
        return None
    return _compact({
        "name": name,
        "start": _date_text(raw.get("start") or raw.get("start_date")),
        "end": _date_text(raw.get("end") or raw.get("end_date")),
        "description": _text(raw.get("description"), 1500),
    })


def _volunteering_item(raw: Any) -> Optional[Dict[str, Any]]:
    if not isinstance(raw, dict):
        return None
    role = _line(raw.get("role") or raw.get("position") or raw.get("title"), 200)
    organization = _org_name(raw.get("company") or raw.get("organization"))
    if not (role or organization):
        return None
    return _compact({
        "role": role,
        "organization": organization,
        "cause": _line(raw.get("cause"), 120),
        "start": _date_text(raw.get("start") or raw.get("start_date")),
        "end": _date_text(raw.get("end") or raw.get("end_date")),
        "description": _text(raw.get("description"), 1000),
    })


def _items(value: Any, build, limit: int) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    for raw in _as_list(value):
        item = build(raw)
        if item:
            out.append(item)
        if len(out) >= limit:
            break
    return out


def _contact_values(value: Any, keys: Tuple[str, ...], limit: int = 3) -> List[str]:
    out: List[str] = []
    for item in _as_list(value):
        if isinstance(item, dict):
            item = next((item.get(k) for k in keys if item.get(k)), "")
        text = _line(item, 200)
        if text and text not in out:
            out.append(text)
        if len(out) >= limit:
            break
    return out


def normalize_linkedin_profile(profile: Any) -> Dict[str, Any]:
    """Compact, JSON-safe copy of a Unipile LinkedIn profile payload.

    Keeps every section a résumé needs -- including what the scoring extraction
    (``_extract_linkedin_profile_data``) drops: role descriptions and locations,
    fields of study, languages, projects, volunteering, websites and any contact
    details LinkedIn shares. Tolerant to the field-name variants Unipile has
    shipped (``position`` / ``title``, ``start`` / ``start_date``, issuer as
    ``organization`` / ``authority``). Idempotent, so it also re-bounds a
    profile that comes back from the browser. Returns ``{}`` for a non-dict.
    """
    if not isinstance(profile, dict):
        return {}
    first = _line(profile.get("first_name") or profile.get("firstName"), 80)
    last = _line(profile.get("last_name") or profile.get("lastName"), 80)
    name = _line(profile.get("name") or profile.get("full_name"), 160) or f"{first} {last}".strip()
    location = profile.get("location") or profile.get("location_name")
    if isinstance(location, dict):
        location = location.get("name") or location.get("default")
    slug = _line(profile.get("public_identifier"), 200).strip("/")
    contact = _as_dict(profile.get("contact_info"))
    websites = _contact_values(profile.get("websites"), ("url", "value"), MAX_LIST)
    return _compact({
        "name": name,
        "first_name": first,
        "last_name": last,
        "headline": _line(profile.get("headline"), 300),
        "location": _line(location, 200),
        "summary": _text(profile.get("summary") or profile.get("about"), MAX_FIELD_CHARS),
        "public_profile_url": public_linkedin_url(
            profile.get("public_profile_url"),
            f"https://www.linkedin.com/in/{slug}" if slug and "/" not in slug and " " not in slug else "",
        ),
        "experience": _items(
            profile.get("work_experience") or profile.get("experience") or profile.get("work_history"),
            _experience_item, MAX_EXPERIENCE,
        ),
        "education": _items(profile.get("education"), _education_item, MAX_EDUCATION),
        "skills": _names(profile.get("skills"), MAX_SKILLS),
        "certifications": _items(
            profile.get("certifications") or profile.get("licenses"), _certification_item, MAX_LIST,
        ),
        "languages": _items(profile.get("languages"), _language_item, MAX_LIST),
        "projects": _items(profile.get("projects"), _project_item, MAX_LIST),
        "volunteering": _items(
            profile.get("volunteering_experience") or profile.get("volunteering"), _volunteering_item, MAX_LIST,
        ),
        "websites": websites,
        # contact_info on a Unipile payload; top-level on an already-normalised one.
        "emails": _contact_values(contact.get("emails") or profile.get("emails"), ("address", "email", "value")),
        "phones": _contact_values(contact.get("phones") or profile.get("phones"), ("number", "phone", "value")),
    })


def deep_search_profile(entry: Any) -> Dict[str, Any]:
    """The part of an Exa deep-search hit that is the person's own profile
    (current title, location, recent roles) -- never the agent's fit rationale."""
    if not isinstance(entry, dict):
        return {}
    return _compact({
        "headline": _line(entry.get("current_title"), 300),
        "location": _line(entry.get("location"), 200),
        "experience": _items(entry.get("recent_companies"), _experience_item, MAX_EXPERIENCE),
    })


# ---------------------------------------------------------------------------
# Names and locations for the JobDiva profile fields
# ---------------------------------------------------------------------------

def is_placeholder_name(name: Any) -> bool:
    """True for names the pipeline invents when it has none ("Unknown Candidate",
    "LinkedIn Professional ab12cd34", a LinkedIn headline, a bare id ...)."""
    text = _line(name, 200)
    key = text.casefold()
    if key in _PLACEHOLDER_NAMES or key.startswith("linkedin professional"):
        return True
    if "|" in text or "@" in text or not re.search(r"[^\W\d_]{2,}", text):
        return True
    # Unipile's last resort before an id: "<title> at <company>". A standalone
    # " at " only -- "Ahmad At-Tamimi" is a name.
    if " at " in f" {key} " and len(text.split()) >= 3:
        return True
    return False


def split_person_name(name: Any) -> Tuple[str, str]:
    """``"Dr. Jane van der Berg, PMP"`` -> ``("Jane", "van der Berg")``.

    Credentials after a comma, honorifics and emoji decorations are dropped;
    the first word is the first name and the rest the last name.
    """
    text = _line(name, 200)
    if "," in text:
        head, tail = text.split(",", 1)
        tail_tokens = [t for t in re.split(r"[\s,]+", tail) if t]
        # "Jane Doe, MBA, PMP" / "Jane Doe, Ph.D." -- credentials, not a name.
        if head.strip() and tail_tokens and all(
            len(t) <= 6 and not t[:1].islower() and sum(ch.isupper() for ch in t) >= 2 for t in tail_tokens
        ):
            text = head
    text = re.sub(r"[^\w\s'.\-]", " ", text, flags=re.UNICODE)
    parts = [p for p in text.split() if p]
    while parts and parts[0].casefold() in _HONORIFICS:
        parts.pop(0)
    if not parts:
        return "", ""
    return parts[0], " ".join(parts[1:])


def jobdiva_address_fields(location: Any) -> Dict[str, str]:
    """UpdateCandidateProfileDef address fields for a free-text location.

    ``"Jersey City, New Jersey, United States"`` -> ``{"city": "Jersey City",
    "state": "NJ", "countryid": "US"}``; ``"Toronto, Ontario, Canada"`` ->
    ``{"city": "Toronto", "state": "ON", "countryid": "CA"}``. Only what the
    string actually states is returned (a bare "United States" gives just the
    country); work arrangements ("Remote", "Hybrid") are never a place.
    """
    from services.location import sanitize_candidate_location
    from services.us_state_index import resolve_state_code

    text = sanitize_candidate_location(_line(location, 200))
    # Exa writes the country with its code: "Jersey City, New Jersey, United States (US)".
    text = re.sub(r"\s*\([A-Za-z]{2,3}\)\s*$", "", text).strip(" ,")
    if not text:
        return {}
    out: Dict[str, str] = {}

    zip_matches = list(re.finditer(r"\b(\d{5})(?:-\d{4})?\b", text))
    if zip_matches:
        out["zipCode"] = zip_matches[-1].group(1)
        text = (text[:zip_matches[-1].start()] + text[zip_matches[-1].end():]).strip(" ,")

    try:  # the Step-5 location parser's alias tables; no instance is needed
        from services.unified_candidate_search import UnifiedCandidateSearch as _U
        country_aliases = dict(_U._COUNTRY_ALIASES)
        province_codes = set(_U._CA_PROVINCE_CODES)
    except Exception:  # pragma: no cover - defensive: keep provisioning alive
        country_aliases = {"US": "US", "USA": "US", "UNITED STATES": "US", "CANADA": "CA", "CA": "CA"}
        province_codes = set()
    provinces = {
        "alberta": "AB", "british columbia": "BC", "manitoba": "MB", "new brunswick": "NB",
        "newfoundland and labrador": "NL", "nova scotia": "NS", "ontario": "ON",
        "prince edward island": "PE", "quebec": "QC", "québec": "QC", "saskatchewan": "SK",
    }

    parts = [p.strip() for p in text.split(",") if p.strip()]
    if len(parts) > 1 and parts[-1].upper().strip(" .") in country_aliases:
        out["countryid"] = country_aliases[parts[-1].upper().strip(" .")]
        parts = parts[:-1]
    elif len(parts) == 1 and parts[0].upper().strip(" .") in country_aliases and not resolve_state_code(parts[0]):
        out["countryid"] = country_aliases[parts[0].upper().strip(" .")]
        return out

    def _region(value: str) -> Tuple[str, str]:
        """(state code, country) for a US state / Canadian province, else ("", "")."""
        us = resolve_state_code(value)
        if us:
            return us.upper(), "US"
        key = value.strip().casefold()
        if key in provinces:
            return provinces[key], "CA"
        # Province codes are disjoint from US state codes ("Toronto, ON").
        if value.strip().upper() in province_codes:
            return value.strip().upper(), "CA"
        return "", ""

    city = ""
    if len(parts) >= 2:
        city = parts[0]
        state, country = _region(parts[1])
        if state:
            out["state"] = state
            out.setdefault("countryid", country)
        elif out.get("countryid") and out["countryid"] != "US":
            out["state"] = parts[1]  # a non-US region, kept as written
    elif len(parts) == 1:
        state, country = _region(parts[0])
        if state:
            out["state"] = state
            out.setdefault("countryid", country)
        else:
            city = parts[0]

    # LinkedIn metro strings: "Greater Chicago Area", "San Francisco Bay Area".
    city = re.sub(r"^greater\s+", "", city, flags=re.IGNORECASE)
    city = re.sub(r"\s+(?:metropolitan|metro|bay)?\s*area$", "", city, flags=re.IGNORECASE).strip()
    if city and re.search(r"[^\W\d_]{2,}", city):
        out["city"] = city
    return out


# ---------------------------------------------------------------------------
# The résumé
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class ProfileResume:
    """A rendered résumé plus what it is made of."""

    text: str
    name: str
    headline: str
    location: str
    # Body sections that carry content: "resume" (the row's own résumé text),
    # "summary", "experience", "education", "skills", "certifications",
    # "languages", "projects", "volunteering", "profile_text".
    sections: Tuple[str, ...]

    @property
    def uses_source_resume(self) -> bool:
        return "resume" in self.sections

    @property
    def is_thin(self) -> bool:
        """No work history, summary or résumé text: re-reading the LinkedIn
        profile is worth it before a JobDiva profile is created from this."""
        return not (_CAREER_SECTIONS & set(self.sections))

    @property
    def is_blank(self) -> bool:
        """Nothing but a name: creating a JobDiva profile from this would make
        exactly the blank profile this module exists to prevent."""
        return not self.sections and not self.headline


def _looks_like_real_resume(text: str) -> bool:
    from services.sourced_candidates_storage import _has_real_resume_text

    return len(text) >= 40 and _has_real_resume_text(text)


def source_resume_text(row: Dict[str, Any]) -> str:
    """The row's own résumé text, when it has a real one.

    LinkedIn rows never do: their ``resume_text`` is Exa's crawl of the profile
    (rendered as profile text instead), the deep-search agent's fit rationale
    (PAIR's words, not the candidate's -- never shown as their résumé), or
    empty. JobDiva / Dice / manual uploads carry real résumés.
    """
    if str(row.get("source") or "").startswith("LinkedIn"):
        return ""
    text = _text(row.get("resume_text"), 60000)
    return text if _looks_like_real_resume(text) else ""


def clean_profile_text(text: Any) -> str:
    """Exa's LinkedIn crawl text, minus its "[...]" separators, markdown and noise."""
    cleaned = str(text or "").replace("\r\n", "\n").replace("\r", "\n")
    cleaned = re.sub(r"^\s*\[\.\.\.\]\s*$", "", cleaned, flags=re.MULTILINE)
    cleaned = cleaned.replace("[...]", " ")
    cleaned = re.sub(r"^\s{0,3}#{1,6}\s*", "", cleaned, flags=re.MULTILINE)
    cleaned = re.sub(r"\*\*(.+?)\*\*", r"\1", cleaned)
    cleaned = _PROFILE_TEXT_NOISE_RE.sub("", cleaned)
    return _text(cleaned, MAX_PROFILE_TEXT_CHARS)


def _profile_text(row: Dict[str, Any], data: Dict[str, Any]) -> str:
    source = str(row.get("source") or "")
    if not source.startswith("LinkedIn") or source == "LinkedIn-DeepSearch":
        return ""
    raw = str(data.get("deep_text") or "").strip() or str(row.get("resume_text") or "")
    text = clean_profile_text(raw)
    return text if len(text) >= 80 else ""


def _coerce_experience(value: Any) -> List[Dict[str, Any]]:
    """company_experience in any stored shape (LinkedIn extraction, LLM parse)."""
    items: List[Dict[str, Any]] = []
    for raw in _as_list(value):
        if not isinstance(raw, dict):
            continue
        item = _experience_item(raw)
        if not item:
            continue
        end_client = _line(raw.get("end_client"), 200)
        if end_client and end_client.casefold() != str(item.get("company", "")).casefold():
            item["company"] = f"{item['company']} (client: {end_client})" if item.get("company") else end_client
        items.append(item)
        if len(items) >= MAX_EXPERIENCE:
            break
    return items


def _skill_names(*values: Any) -> List[str]:
    out: List[str] = []
    seen = set()
    for value in values:
        for name in _names(value, MAX_SKILLS):
            if name.casefold() not in seen:
                seen.add(name.casefold())
                out.append(name)
    return out[:MAX_SKILLS]


def _render_experience(items: List[Dict[str, Any]]) -> List[str]:
    lines: List[str] = []
    for item in items:
        title, company = item.get("title", ""), item.get("company", "")
        lines.append(f"{title} at {company}" if title and company else (title or company))
        meta = " · ".join(p for p in (
            _date_range(item.get("start", ""), item.get("end", ""), bool(item.get("current"))),
            item.get("location", ""),
        ) if p)
        if meta:
            lines.append(meta)
        if item.get("description"):
            lines.append(item["description"])
        if item.get("skills"):
            lines.append("Skills: " + ", ".join(item["skills"]))
        lines.append("")
    return lines


def _render_education(items: List[Dict[str, Any]]) -> List[str]:
    lines: List[str] = []
    for item in items:
        degree = ", ".join(p for p in (item.get("degree", ""), item.get("field_of_study", "")) if p)
        lines.append(item.get("school") or degree)
        meta = " · ".join(p for p in (
            degree if item.get("school") else "",
            " - ".join(p for p in (item.get("start", ""), item.get("end", "")) if p),
        ) if p)
        if meta:
            lines.append(meta)
        if item.get("description"):
            lines.append(item["description"])
        lines.append("")
    return lines


def _render_simple(items: List[Dict[str, Any]], *keys: str) -> List[str]:
    return ["- " + " · ".join(str(item[k]) for k in keys if item.get(k)) for item in items]


def _render_projects(items: List[Dict[str, Any]], heading_keys: Tuple[str, ...]) -> List[str]:
    lines: List[str] = []
    for item in items:
        head = " at ".join(str(item[k]) for k in heading_keys if item.get(k))
        dates = " - ".join(p for p in (item.get("start", ""), item.get("end", "")) if p)
        lines.append(" · ".join(p for p in (head, item.get("cause", ""), dates) if p))
        if item.get("description"):
            lines.append(item["description"])
        lines.append("")
    return lines


def build_profile_resume(
    row: Dict[str, Any],
    data: Optional[Dict[str, Any]] = None,
    *,
    email: str = "",
    phone: str = "",
) -> ProfileResume:
    """Render the résumé JobDiva should get for a ``sourced_candidates`` row.

    ``row`` needs ``name`` / ``headline`` / ``location`` / ``profile_url`` /
    ``resume_text`` / ``source``; ``data`` is its (parsed) data blob. ``email`` /
    ``phone`` are the contact details the caller will write to the profile --
    pass real ones only (a synthetic placeholder has no place on a résumé).

    A row with a real résumé keeps it verbatim under a contact header. Anything
    else -- every LinkedIn row -- gets a LinkedIn-style résumé from, in order of
    preference, the full LinkedIn profile captured at sourcing
    (``data.linkedin_profile``), the structured fields the save payload carries
    (``company_experience`` / ``education`` / ``certifications`` / ``skills``) and
    the LLM extraction (``enhanced_info``), plus Exa's crawl of the profile.
    """
    data = data if isinstance(data, dict) else _as_dict(data if data is not None else row.get("data"))
    profile = _as_dict(data.get("linkedin_profile"))
    enhanced = _as_dict(data.get("enhanced_info")) or _as_dict(row.get("enhanced_info"))
    from services.location import sanitize_candidate_location

    name = _line(row.get("name"), 160)
    if is_placeholder_name(name):
        name = _line(profile.get("name"), 160) or _line(enhanced.get("candidate_name"), 160) or name
    headline = (
        _line(profile.get("headline"), 300)
        or _line(row.get("headline"), 300)
        or _line(enhanced.get("job_title"), 300)
    )
    location = ""
    for candidate_location in (row.get("location"), profile.get("location"), enhanced.get("current_location")):
        location = sanitize_candidate_location(_line(candidate_location, 200))
        if location:
            break
    urls = _as_dict(data.get("urls")) or _as_dict(enhanced.get("urls"))
    linkedin = public_linkedin_url(
        row.get("profile_url"), profile.get("public_profile_url"), urls.get("linkedin"), urls.get("linkedin_url"),
    )
    links = [u for u in [*_as_list(profile.get("websites")), urls.get("github"), urls.get("portfolio")] if _line(u, 400)]

    header = [name] if name else []
    if headline and headline.casefold() == name.casefold():
        headline = ""  # some rows carry the name in the headline column
    header += [line for line in (headline, location) if line]
    if email:
        header.append(f"Email: {email}")
    if phone:
        header.append(f"Phone: {phone}")
    if linkedin:
        header.append(f"LinkedIn: {linkedin}")
    for link in list(dict.fromkeys(_line(u, 400) for u in links))[:3]:
        if link != linkedin:
            header.append(f"Website: {link}")

    sections: List[str] = []
    body: List[str] = []

    def _section(key: str, title: str, lines: Iterable[str]) -> None:
        content = [line for line in lines]
        while content and not content[-1]:
            content.pop()
        if any(line.strip() for line in content):
            sections.append(key)
            body.extend(["", title, *content])

    own_resume = source_resume_text(row)
    if own_resume:
        sections.append("resume")
        body.extend(["", own_resume])
    else:
        summary = _text(profile.get("summary") or data.get("summary"), MAX_FIELD_CHARS)
        experience = (
            _as_list(profile.get("experience"))
            or _coerce_experience(data.get("company_experience"))
            or _coerce_experience(enhanced.get("company_experience"))
        )
        education = (
            _items(profile.get("education"), _education_item, MAX_EDUCATION)
            or _items(data.get("education"), _education_item, MAX_EDUCATION)
            or _items(enhanced.get("candidate_education"), _education_item, MAX_EDUCATION)
        )
        certifications = (
            _items(profile.get("certifications"), _certification_item, MAX_LIST)
            or _items(data.get("certifications"), _certification_item, MAX_LIST)
            or _items(enhanced.get("candidate_certification"), _certification_item, MAX_LIST)
        )
        skills = _skill_names(profile.get("skills"), data.get("skills"), enhanced.get("structured_skills"))

        _section("summary", SECTION_SUMMARY, [summary] if summary else [])
        _section("experience", SECTION_EXPERIENCE, _render_experience(
            [e for e in experience if isinstance(e, dict)][:MAX_EXPERIENCE]
        ))
        _section("education", SECTION_EDUCATION, _render_education(education))
        _section("skills", SECTION_SKILLS, [", ".join(skills)] if skills else [])
        _section("certifications", SECTION_CERTIFICATIONS, _render_simple(certifications, "name", "issuer", "date"))
        _section("languages", SECTION_LANGUAGES, _render_simple(
            _items(profile.get("languages"), _language_item, MAX_LIST), "name", "proficiency",
        ))
        _section("projects", SECTION_PROJECTS, _render_projects(
            _items(profile.get("projects"), _project_item, MAX_LIST), ("name",),
        ))
        _section("volunteering", SECTION_VOLUNTEERING, _render_projects(
            _items(profile.get("volunteering"), _volunteering_item, MAX_LIST), ("role", "organization"),
        ))
        # Exa's crawl often holds the role descriptions the structured parse
        # lacks. The full Unipile profile already has them, so it is skipped then.
        if not profile.get("experience"):
            profile_text = _profile_text(row, data)
            _section("profile_text", SECTION_PROFILE_TEXT, [profile_text] if profile_text else [])

    text = "\n".join([*header, *body]).strip() + "\n"
    return ProfileResume(
        text=text, name=name, headline=headline, location=location, sections=tuple(sections),
    )


# ---------------------------------------------------------------------------
# The Word document
# ---------------------------------------------------------------------------

def _xml_safe(value: str) -> str:
    return _XML_INVALID_RE.sub("", value)


def resume_to_docx(resume: ProfileResume) -> Optional[bytes]:
    """The résumé as a .docx (bytes), or None when python-docx is unavailable or
    the build fails -- callers then upload the plain text as a .txt file."""
    try:
        from docx import Document
        from docx.shared import Pt
    except Exception:
        logger.warning("python-docx unavailable; résumé will be uploaded as .txt")
        return None
    try:
        document = Document()
        normal = document.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(2)
        lines = resume.text.split("\n")
        seen_name = False
        previous_blank = True
        for raw in lines:
            line = _xml_safe(raw.rstrip())
            stripped = line.strip()
            if not seen_name:
                if not stripped:
                    continue
                run = document.add_paragraph().add_run(stripped)
                run.bold = True
                run.font.size = Pt(16)
                seen_name = True
                previous_blank = False
                continue
            if not stripped:
                if not previous_blank:
                    document.add_paragraph()
                previous_blank = True
                continue
            previous_blank = False
            if stripped in SECTION_TITLES:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(8)
                run = paragraph.add_run(stripped)
                run.bold = True
                run.font.size = Pt(12)
            elif stripped.startswith(("- ", "• ")):
                document.add_paragraph(stripped[2:].strip(), style="List Bullet")
            else:
                document.add_paragraph(stripped)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()
    except Exception as exc:
        logger.warning("résumé .docx build failed (%s); uploading as .txt", exc)
        return None
